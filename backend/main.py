"""
FastAPI backend para gestionar registros desde un Excel local.

El frontend mantiene los mismos endpoints: listar registros, descargar Excel y
cargar un Excel con diferencias para actualizar la base local.
"""

from pathlib import Path
from typing import Any, Dict, List
import io
import logging
import os
import re

from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
import pandas as pd

load_dotenv()

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parent

app = FastAPI(
    title="Gestion de Eventos API",
    description="API para gestionar eventos e inscripciones desde Excel local",
    version="4.0.0",
)

origins = [
    "http://localhost:4200",
    "http://127.0.0.1:4200",
    "http://localhost:3000",
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def _excel_path() -> Path:
    configured_path = os.getenv("LOCAL_EXCEL_PATH", "../database.xlsx").strip() or "../database.xlsx"
    path = Path(configured_path)
    if not path.is_absolute():
        path = BASE_DIR / path
    return path.resolve()


def _configured_sheet_name() -> str:
    return os.getenv("EXCEL_SHEET_NAME", "").strip()


def _sheet_to_read() -> str | int:
    return _configured_sheet_name() or 0


def _sheet_to_write(path: Path) -> str:
    configured_sheet_name = _configured_sheet_name()
    if configured_sheet_name:
        return configured_sheet_name[:31]

    if path.exists():
        workbook = load_workbook(path, read_only=True)
        try:
            return workbook.sheetnames[0][:31] if workbook.sheetnames else "Registros"
        finally:
            workbook.close()

    return "Registros"


def _normalizar_df(df: pd.DataFrame, columnas: List[str]) -> List[Dict[str, Any]]:
    resultado: List[Dict[str, Any]] = []
    for _, row in df.iterrows():
        registro: Dict[str, Any] = {}
        for col in columnas:
            valor = row.get(col)
            registro[col] = "" if pd.isna(valor) else str(valor).strip()
        resultado.append(registro)
    return resultado


def _normalizar_columnas(columnas: List[str]) -> List[str]:
    resultado = []
    for col in columnas:
        col_normalizado = str(col)
        for special_space in [
            "\xa0",
            "\u2000",
            "\u2001",
            "\u2002",
            "\u2003",
            "\u2004",
            "\u2005",
            "\u2006",
            "\u2007",
            "\u2008",
            "\u2009",
            "\u200a",
        ]:
            col_normalizado = col_normalizado.replace(special_space, " ")
        col_normalizado = re.sub(r"\s+", " ", col_normalizado).strip()
        resultado.append(col_normalizado)
    return resultado


def _texto_comparable(valor: Any) -> str:
    if valor is None or pd.isna(valor):
        return ""
    texto = str(valor).strip().lower()
    texto = re.sub(r"\s+", " ", texto)
    return texto


def _fecha_comparable(valor: Any) -> str:
    if valor is None or pd.isna(valor):
        return ""

    texto = str(valor).strip()
    if not texto:
        return ""

    fecha = pd.to_datetime(valor, errors="coerce", dayfirst=True)
    if pd.isna(fecha):
        fecha = pd.to_datetime(texto, errors="coerce")
    if pd.isna(fecha):
        return texto.split()[0].replace("/", "-").strip()

    return fecha.strftime("%d-%m-%Y")


def _buscar_columna(columnas: List[str], patrones: List[str]) -> str | None:
    columnas_norm = {col: _texto_comparable(col) for col in columnas}
    for patron in patrones:
        patron_norm = _texto_comparable(patron)
        for col, col_norm in columnas_norm.items():
            if patron_norm in col_norm:
                return col
    return None


def _columnas_duplicado(columnas: List[str]) -> Dict[str, str]:
    completion_col = _buscar_columna(columnas, ["completion time"])
    curso_col = _buscar_columna(columnas, ["en que curso", "en que curso, taller", "curso, taller", "curso"])
    matricula_col = _buscar_columna(columnas, ["matricula", "matrícula", "nomina", "nómina"])
    nombre_col = _buscar_columna(columnas, ["nombre completo", "name"])

    faltantes = []
    if not completion_col:
        faltantes.append("Completion time")
    if not curso_col:
        faltantes.append("En que curso")
    if not matricula_col and not nombre_col:
        faltantes.append("Nombre o matricula")

    if faltantes:
        raise HTTPException(
            status_code=400,
            detail=(
                "No se pudieron encontrar las columnas necesarias para validar duplicados: "
                + ", ".join(faltantes)
            ),
        )

    return {
        "completion": completion_col,
        "curso": curso_col,
        "matricula": matricula_col or "",
        "nombre": nombre_col or "",
    }


def _datos_duplicado(row: Dict[str, Any], columnas_dup: Dict[str, str]) -> Dict[str, str]:
    return {
        "fecha": _fecha_comparable(row.get(columnas_dup["completion"])),
        "curso": _texto_comparable(row.get(columnas_dup["curso"])),
        "matricula": _texto_comparable(row.get(columnas_dup["matricula"])) if columnas_dup["matricula"] else "",
        "nombre": _texto_comparable(row.get(columnas_dup["nombre"])) if columnas_dup["nombre"] else "",
    }


def _es_misma_persona(a: Dict[str, str], b: Dict[str, str]) -> bool:
    misma_matricula = bool(a["matricula"] and b["matricula"] and a["matricula"] == b["matricula"])
    mismo_nombre = bool(a["nombre"] and b["nombre"] and a["nombre"] == b["nombre"])
    return misma_matricula or mismo_nombre


def _es_registro_duplicado(a: Dict[str, str], b: Dict[str, str]) -> bool:
    return bool(
        a["fecha"]
        and b["fecha"]
        and a["fecha"] == b["fecha"]
        and a["curso"]
        and b["curso"]
        and a["curso"] == b["curso"]
        and _es_misma_persona(a, b)
    )


def _buscar_duplicado(
    row: Dict[str, Any],
    registros: List[Dict[str, Any]],
    columnas_dup: Dict[str, str],
) -> Dict[str, Any] | None:
    datos_row = _datos_duplicado(row, columnas_dup)
    for registro in registros:
        datos_registro = _datos_duplicado(registro, columnas_dup)
        if _es_registro_duplicado(datos_row, datos_registro):
            return registro
    return None


def _leer_excel_local() -> pd.DataFrame:
    path = _excel_path()
    if not path.exists():
        raise HTTPException(
            status_code=500,
            detail=f"No existe el archivo Excel local: {path}",
        )

    logger.info("Leyendo Excel local: %s", path)
    configured_sheet_name = _configured_sheet_name()
    try:
        df = pd.read_excel(path, sheet_name=_sheet_to_read())
    except ValueError as e:
        if configured_sheet_name and "Worksheet named" in str(e):
            logger.warning("No se encontro la hoja '%s'. Leyendo la primera hoja.", configured_sheet_name)
            df = pd.read_excel(path, sheet_name=0)
        else:
            raise

    df.columns = [str(col).strip() for col in df.columns.tolist()]
    return df


def _guardar_excel_local(registros: List[Dict[str, Any]], columnas: List[str]) -> None:
    path = _excel_path()
    path.parent.mkdir(parents=True, exist_ok=True)

    df = pd.DataFrame(registros, columns=columnas)
    sheet_name = _sheet_to_write(path)
    logger.info("Guardando %s registros en Excel local: %s", len(df), path)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)


def _validar_archivo(file: UploadFile) -> None:
    if not file.filename:
        raise HTTPException(status_code=400, detail="No se recibio archivo.")

    nombre = file.filename.lower()
    if not (nombre.endswith(".xlsx") or nombre.endswith(".xls")):
        raise HTTPException(status_code=400, detail="El archivo debe ser Excel (.xlsx o .xls).")


def _leer_excel_subido(contenido: bytes) -> pd.DataFrame:
    configured_sheet_name = _configured_sheet_name()
    try:
        return pd.read_excel(io.BytesIO(contenido), sheet_name=_sheet_to_read())
    except ValueError as e:
        if configured_sheet_name and "Worksheet named" in str(e):
            return pd.read_excel(io.BytesIO(contenido), sheet_name=0)
        raise


@app.get("/")
async def root():
    return {
        "message": "Gestion de Eventos API",
        "version": "4.0.0",
        "storage": "local_excel",
        "docs": "/docs",
    }


@app.get("/health")
async def health_check():
    try:
        df = _leer_excel_local()
        return {
            "status": "healthy",
            "storage": "local_excel",
            "path": str(_excel_path()),
            "rows": len(df),
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Excel local no disponible: {str(e)}")


@app.get("/api/datos")
async def obtener_datos():
    try:
        df = _leer_excel_local()
        columnas = df.columns.tolist()
        datos = _normalizar_df(df, columnas)
        return {
            "success": True,
            "total": len(datos),
            "columns": columnas,
            "data": datos,
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error("Error al obtener datos del Excel local: %s", str(e))
        raise HTTPException(status_code=500, detail=f"Error al obtener datos: {str(e)}")


@app.get("/api/datos/excel")
async def descargar_datos_excel():
    try:
        df = _leer_excel_local()
        output = io.BytesIO()
        sheet_name = _sheet_to_write(_excel_path())
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="registros_excel_local.xlsx"'},
        )
    except Exception as e:
        logger.error("Error al exportar datos: %s", str(e))
        raise HTTPException(status_code=500, detail=f"Error al exportar datos: {str(e)}")


@app.post("/api/datos/upload-excel")
async def cargar_excel_diferencial(file: UploadFile = File(...)):
    try:
        _validar_archivo(file)

        contenido = await file.read()
        if not contenido:
            raise HTTPException(status_code=400, detail="El archivo Excel esta vacio.")

        df_base = _leer_excel_local()
        columnas_base = _normalizar_columnas(df_base.columns.tolist())
        df_base.columns = columnas_base
        base_norm = _normalizar_df(df_base, columnas_base)

        df_subido = _leer_excel_subido(contenido)
        columnas_subidas = _normalizar_columnas(df_subido.columns.tolist())
        df_subido.columns = columnas_subidas

        if columnas_subidas != columnas_base:
            raise HTTPException(
                status_code=400,
                detail=(
                    "El formato del archivo no coincide con el Excel base. "
                    f"Columnas esperadas: {columnas_base}. "
                    f"Columnas recibidas: {columnas_subidas}."
                ),
            )

        subido_norm = _normalizar_df(df_subido, columnas_base)
        columnas_dup = _columnas_duplicado(columnas_base)
        datos_actualizados = list(base_norm)
        insertados = 0
        duplicados = 0

        for row_subido in subido_norm:
            if _buscar_duplicado(row_subido, datos_actualizados, columnas_dup) is None:
                datos_actualizados.append(row_subido)
                insertados += 1
            else:
                duplicados += 1

        if insertados > 0:
            _guardar_excel_local(datos_actualizados, columnas_base)

        return {
            "success": True,
            "message": (
                "Cambios guardados en el Excel local."
                if insertados > 0
                else "No se detectaron registros nuevos respecto al Excel base."
            ),
            "columns": columnas_base,
            "summary": {
                "total_base": len(base_norm),
                "total_subido": len(subido_norm),
                "insertados": insertados,
                "actualizados": 0,
                "sin_cambios": duplicados,
                "duplicados": duplicados,
                "diferentes": insertados,
                "criterio_duplicado": "misma persona por nombre o matricula + mismo curso + misma fecha de Completion time",
            },
            "data": datos_actualizados,
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error("Error al cargar Excel diferencial: %s", str(e))
        raise HTTPException(status_code=400, detail=f"Error al procesar archivo: {str(e)}")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
